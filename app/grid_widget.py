"""Interactive sheet grid with double-click editing (streamlit-aggrid)."""

from typing import Any

import pandas as pd
from openpyxl.utils import get_column_letter
from st_aggrid import AgGrid, GridOptionsBuilder, JsCode
from st_aggrid.shared import DataReturnMode, GridUpdateMode

from app.cell_store import CellStore

_CHAR_PX = 7
_MIN_COL_WIDTH = 36
_MAX_COL_WIDTH = 480
_ROW_COL_WIDTH = 48


def _column_width(df: pd.DataFrame, col: str, bundled_px: int | None) -> tuple[int, bool]:
    max_len = len(col)
    if col in df.columns:
        for val in df[col].fillna("").astype(str):
            max_len = max(max_len, len(val))

    if max_len > 50:
        width = min(max(bundled_px or 140, 120), 320)
        return width, True

    content_px = min(max(max_len * _CHAR_PX + 20, _MIN_COL_WIDTH), _MAX_COL_WIDTH)
    if bundled_px:
        return min(max(bundled_px, content_px), _MAX_COL_WIDTH), False
    return content_px, False


def build_sheet_dataframe(
    store: CellStore, sheet: str
) -> tuple[pd.DataFrame, list[str], dict[str, Any]]:
    meta = store.sheet_meta(sheet)
    bounds = meta.get("bounds")
    if not bounds:
        return pd.DataFrame(), [], {}

    min_r, max_r, min_c, max_c = bounds
    columns = [get_column_letter(c) for c in range(min_c, max_c + 1)]
    bundled_widths = meta.get("col_widths", {})
    rows: list[dict[str, Any]] = []

    for r in range(min_r, max_r + 1):
        row: dict[str, Any] = {"Row": r}
        for c in range(min_c, max_c + 1):
            letter = get_column_letter(c)
            coord = f"{letter}{r}"
            row[letter] = store.display_value(sheet, coord)
            row[f"__e_{letter}"] = store.is_editable(sheet, coord)
        rows.append(row)

    return pd.DataFrame(rows), columns, {
        "min_r": min_r,
        "min_c": min_c,
        "sheet": sheet,
        "bundled_widths": bundled_widths,
    }


_EDITABLE_JS = JsCode(
    """
    function(params) {
        const field = params.colDef.field;
        if (!field || field === 'Row') return false;
        const key = '__e_' + field;
        return params.data[key] === true;
    }
    """
)

_AUTOSIZE_JS = JsCode(
    """
    function(params) {
        try {
            if (params.api && params.api.autoSizeAllColumns) {
                params.api.autoSizeAllColumns(false);
            } else if (params.columnApi) {
                params.columnApi.autoSizeAllColumns(false);
            }
        } catch (e) {}
    }
    """
)


def _rows_as_dicts(data: Any) -> list[dict[str, Any]]:
    if data is None:
        return []
    if isinstance(data, pd.DataFrame):
        return data.to_dict("records")
    if isinstance(data, list):
        return [row for row in data if isinstance(row, dict)]
    return []


def render_editable_grid(
    store: CellStore,
    sheet: str,
    height: int = 680,
) -> dict[str, str] | None:
    df, data_columns, ctx = build_sheet_dataframe(store, sheet)
    if df.empty:
        return None

    bundled = ctx.get("bundled_widths", {})
    min_c = ctx["min_c"]

    gb = GridOptionsBuilder.from_dataframe(df)
    gb.configure_default_column(
        editable=False,
        filterable=False,
        sortable=False,
        resizable=True,
        wrapText=False,
        autoHeight=False,
        singleClickEdit=False,
        suppressSizeToFit=True,
    )
    gb.configure_column(
        "Row",
        editable=False,
        width=_ROW_COL_WIDTH,
        minWidth=_ROW_COL_WIDTH,
        maxWidth=_ROW_COL_WIDTH,
        pinned="left",
        cellStyle={"backgroundColor": "#f0f0f0"},
        suppressSizeToFit=True,
    )

    for idx, col in enumerate(data_columns):
        col_idx = min_c + idx
        bundled_px = bundled.get(str(col_idx)) or bundled.get(col_idx)
        width, wrap = _column_width(df, col, bundled_px)
        gb.configure_column(
            col,
            editable=_EDITABLE_JS,
            width=width,
            minWidth=min(width, _MIN_COL_WIDTH),
            maxWidth=_MAX_COL_WIDTH,
            wrapText=wrap,
            autoHeight=wrap,
            suppressSizeToFit=True,
            cellStyle=JsCode(
                f"""
                function(params) {{
                    if (params.data['__e_{col}']) {{
                        return {{backgroundColor: '#ffffff'}};
                    }}
                    return {{backgroundColor: '#f3f3f3'}};
                }}
                """
            ),
        )
    for col in data_columns:
        gb.configure_column(f"__e_{col}", hide=True)

    gb.configure_grid_options(
        domLayout="normal",
        suppressMovableColumns=True,
        enterNavigatesVertically=True,
        enterNavigatesVerticallyAfterEdit=True,
        stopEditingWhenCellsLoseFocus=True,
        onFirstDataRendered=_AUTOSIZE_JS,
    )

    grid_options = gb.build()
    response = AgGrid(
        df,
        gridOptions=grid_options,
        height=height,
        width="100%",
        update_mode=GridUpdateMode.VALUE_CHANGED,
        data_return_mode=DataReturnMode.AS_INPUT,
        allow_unsafe_jscode=True,
        theme="streamlit",
        fit_columns_on_grid_load=False,
        key=f"aggrid_{ctx['sheet']}",
    )

    if response is None:
        return None

    updated = response.get("data") if hasattr(response, "get") else getattr(response, "data", None)
    rows = _rows_as_dicts(updated)
    if not rows:
        return None

    edits: dict[str, str] = {}
    for row_data in rows:
        r = row_data.get("Row")
        if r is None:
            continue
        for col in data_columns:
            coord = f"{col}{r}"
            if not store.is_editable(sheet, coord):
                continue
            new_val = row_data.get(col, "")
            old_val = store.display_value(sheet, coord)
            new_str = "" if new_val is None else str(new_val)
            if new_str != old_val:
                edits[coord] = new_str

    return edits if edits else None
