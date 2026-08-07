from html import escape
from typing import Any

from openpyxl.utils import get_column_letter

from app.conditional_formatting import apply_conditional_formatting, rating_color
from app.utils import color_to_css, column_width_px, format_display_value, row_height_px
from app.workbook_manager import WorkbookManager


def _cell_style(ws, cell, display_value: Any) -> dict[str, str]:
    css: dict[str, str] = {
        "border": "1px solid #b4b4b4",
        "padding": "1px 4px",
        "font-family": "Calibri, Arial, sans-serif",
        "font-size": "11px",
        "vertical-align": "middle",
        "white-space": "nowrap",
        "overflow": "hidden",
        "text-overflow": "ellipsis",
        "max-width": "220px",
    }
    if cell.font:
        if cell.font.bold:
            css["font-weight"] = "bold"
        if cell.font.size:
            css["font-size"] = f"{cell.font.size}pt"
        fg = color_to_css(cell.font.color)
        if fg:
            css["color"] = fg
    if cell.fill and cell.fill.patternType and cell.fill.patternType != "none":
        bg = color_to_css(cell.fill.fgColor)
        if bg:
            css["background-color"] = bg
    if cell.alignment:
        if cell.alignment.horizontal:
            css["text-align"] = cell.alignment.horizontal
        if cell.alignment.vertical:
            css["vertical-align"] = cell.alignment.vertical

    rating_bg = rating_color(display_value)
    if rating_bg and "background-color" not in css:
        css["background-color"] = rating_bg

    return apply_conditional_formatting(ws, cell.coordinate, display_value, css)


def _css_dict_to_str(css: dict[str, str]) -> str:
    return "; ".join(f"{k}: {v}" for k, v in css.items())


def _merged_map(ws) -> tuple[dict[tuple[int, int], tuple[int, int, int, int]], set[tuple[int, int]]]:
    merged: dict[tuple[int, int], tuple[int, int, int, int]] = {}
    covered: set[tuple[int, int]] = set()
    for mrange in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = mrange.bounds
        merged[(min_row, min_col)] = (min_row, min_col, max_row, max_col)
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                if (r, c) != (min_row, min_col):
                    covered.add((r, c))
    return merged, covered


def render_sheet_html(manager: WorkbookManager, sheet: str) -> str:
    ws = manager.wb[sheet]
    bounds = manager.sheet_bounds(sheet)
    if bounds is None:
        return "<p><em>This sheet is empty. Use it as in Excel to add content.</em></p>"

    min_r, max_r, min_c, max_c = bounds
    merged, covered = _merged_map(ws)

    colgroup = "".join(
        f"<col style='width:{column_width_px(ws, c)}px'>" for c in range(min_c, max_c + 1)
    )

    header_cells = "".join(
        f"<th style='background:#217346;color:white;font-size:10px;padding:2px;border:1px solid #999'>"
        f"{get_column_letter(c)}</th>"
        for c in range(min_c, max_c + 1)
    )

    body_rows: list[str] = []
    for r in range(min_r, max_r + 1):
        height = row_height_px(ws, r)
        cells: list[str] = []
        for c in range(min_c, max_c + 1):
            if (r, c) in covered:
                continue
            coord = f"{get_column_letter(c)}{r}"
            cell = ws[coord]
            display = manager.get_display_value(sheet, coord)
            text = escape(format_display_value(display))
            css = _cell_style(ws, cell, display)

            if manager.is_formula_cell(sheet, coord):
                css.setdefault("background-color", "#f3f3f3")

            rowspan, colspan = 1, 1
            if (r, c) in merged:
                min_row, min_col, max_row, max_col = merged[(r, c)]
                rowspan = max_row - min_row + 1
                colspan = max_col - min_col + 1

            attrs = [f"style='{_css_dict_to_str(css)}'", f"title='{coord}'"]
            if rowspan > 1:
                attrs.append(f"rowspan='{rowspan}'")
            if colspan > 1:
                attrs.append(f"colspan='{colspan}'")
            cells.append(f"<td {' '.join(attrs)}>{text}</td>")

        body_rows.append(
            f"<tr style='height:{height}px'>"
            f"<th style='background:#f0f0f0;color:#333;font-size:10px;padding:2px;border:1px solid #b4b4b4;width:36px'>{r}</th>"
            f"{''.join(cells)}"
            "</tr>"
        )

    return f"""
    <div class='excel-sheet-wrap'>
      <table class='excel-sheet'>
        <colgroup>
          <col style='width:36px'>
          {colgroup}
        </colgroup>
        <thead>
          <tr>
            <th style='background:#217346;color:white;font-size:10px;padding:2px;border:1px solid #999'></th>
            {header_cells}
          </tr>
        </thead>
        <tbody>
          {''.join(body_rows)}
        </tbody>
      </table>
    </div>
    """


def render_sheet_css() -> str:
    return """
    <style>
      .excel-sheet-wrap {
        overflow: auto;
        max-height: 78vh;
        border: 1px solid #a6a6a6;
        background: #e6e6e6;
        padding: 8px;
      }
      .excel-sheet {
        border-collapse: collapse;
        background: white;
        table-layout: fixed;
      }
      .excel-sheet td, .excel-sheet th {
        box-sizing: border-box;
      }
    </style>
    """
