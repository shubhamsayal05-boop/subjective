#!/usr/bin/env python3
"""One-time extractor: dump xlsm templates to Python JSON bundles."""

import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from app.utils import color_to_css, column_width_px, format_display_value, is_formula, sheet_bounds

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "templates"
OUT = ROOT / "app" / "bundled"


def extract_cell_style(cell) -> dict:
    style: dict = {}
    if cell.font and cell.font.bold:
        style["bold"] = True
    if cell.font and cell.font.color:
        fg = color_to_css(cell.font.color)
        if fg:
            style["fg"] = fg
    if cell.fill and cell.fill.patternType and cell.fill.patternType != "none":
        bg = color_to_css(cell.fill.fgColor)
        if bg:
            style["bg"] = bg
    return style


def extract_workbook(path: Path) -> dict:
    wb = load_workbook(path, keep_vba=False, data_only=False)
    wb_vals = load_workbook(path, data_only=True)

    named_ranges: dict[str, str] = {}
    for name, dn in wb.defined_names.items():
        dest = list(dn.destinations)
        if dest:
            sheet, ref = dest[0]
            named_ranges[name] = f"{sheet}!{ref}"

    sheets: dict = {}
    for ws in wb.worksheets:
        bounds = sheet_bounds(ws)
        if bounds is None:
            sheets[ws.title] = {
                "bounds": None,
                "merges": [],
                "col_widths": {},
                "cells": {},
            }
            continue

        min_r, max_r, min_c, max_c = bounds
        ws_vals = wb_vals[ws.title]
        cells: dict = {}
        for r in range(min_r, max_r + 1):
            for c in range(min_c, max_c + 1):
                from openpyxl.utils import get_column_letter

                coord = f"{get_column_letter(c)}{r}"
                cell = ws[coord]
                raw = cell.value
                cached = ws_vals[coord].value
                entry: dict = {}
                if is_formula(raw):
                    entry["formula"] = raw
                    if cached is not None:
                        entry["cached"] = format_display_value(cached)
                elif raw is not None:
                    entry["value"] = format_display_value(raw)

                style = extract_cell_style(cell)
                if style:
                    entry["style"] = style
                if entry:
                    cells[coord] = entry

        merges = []
        for mrange in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = mrange.bounds
            merges.append(
                {
                    "r": min_row,
                    "c": min_col,
                    "rs": max_row - min_row + 1,
                    "cs": max_col - min_col + 1,
                }
            )

        col_widths = {
            c: max(column_width_px(ws, c), 28) for c in range(min_c, max_c + 1)
        }

        sheets[ws.title] = {
            "bounds": [min_r, max_r, min_c, max_c],
            "merges": merges,
            "col_widths": col_widths,
            "cells": cells,
        }

    return {
        "book_name": path.name,
        "named_ranges": named_ranges,
        "sheets": sheets,
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    mapping = {
        "base": "Subjective_SprdSheet_072926.xlsm",
        "bev": "BEV_Subjective_SprdSheet_072926.xlsm",
        "cvt": "CVT Subjective_SprdSheet_072926.xlsm",
    }
    for key, filename in mapping.items():
        path = TEMPLATES / filename
        if not path.exists():
            print(f"Skip missing {path}")
            continue
        print(f"Extracting {filename}...")
        data = extract_workbook(path)
        out_path = OUT / f"{key}.json"
        with open(out_path, "w", encoding="utf-8") as f:
            json.dump(data, f, separators=(",", ":"))
        print(f"  -> {out_path} ({out_path.stat().st_size / 1024 / 1024:.1f} MB)")


if __name__ == "__main__":
    main()
