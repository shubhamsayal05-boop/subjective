from openpyxl.utils import get_column_letter

from app.conditional_formatting import rating_color
from app.utils import color_to_css, column_width_px, format_display_value, sheet_bounds
from app.workbook_manager import WorkbookManager


def _cell_bg(ws, cell, display_value) -> str | None:
    if cell.fill and cell.fill.patternType and cell.fill.patternType != "none":
        bg = color_to_css(cell.fill.fgColor)
        if bg:
            return bg
    return rating_color(display_value)


def build_grid_data(manager: WorkbookManager, sheet: str, height: int = 700) -> dict:
    ws = manager.wb[sheet]
    bounds = sheet_bounds(ws)
    if bounds is None:
        return {
            "sheet": sheet,
            "minRow": 1,
            "maxRow": 1,
            "minCol": 1,
            "maxCol": 1,
            "cells": {},
            "merges": [],
            "colWidths": {},
            "height": height,
        }

    min_r, max_r, min_c, max_c = bounds
    cells: dict[str, dict] = {}
    for r in range(min_r, max_r + 1):
        for c in range(min_c, max_c + 1):
            coord = f"{get_column_letter(c)}{r}"
            cell = ws[coord]
            display = manager.get_display_value(sheet, coord)
            info: dict = {
                "v": format_display_value(display),
                "e": manager.is_editable(sheet, coord),
            }
            bg = _cell_bg(ws, cell, display)
            if bg:
                info["bg"] = bg
            cells[coord] = info

    merges = []
    for mrange in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = mrange.bounds
        merges.append(
            {
                "r": min_row,
                "c": min_col,
                "rs": max_row - min_row + 1,
                "cs": max_col - min_col + 1,
            }
        )

    col_widths = {c: max(column_width_px(ws, c), 28) for c in range(min_c, max_c + 1)}

    return {
        "sheet": sheet,
        "minRow": min_r,
        "maxRow": max_r,
        "minCol": min_c,
        "maxCol": max_c,
        "cells": cells,
        "merges": merges,
        "colWidths": col_widths,
        "height": height,
    }
