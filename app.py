import streamlit as st
import formulas

from app.config import get_template_path, ordered_sheets, VERSIONS
from app.html_renderer import render_sheet_css, render_sheet_html
from app.utils import format_display_value
from app.workbook_manager import WorkbookManager


@st.cache_resource(show_spinner="Loading formula engine for {version_label}...")
def load_formula_model(template_path: str, version_label: str) -> formulas.ExcelModel:
    return formulas.ExcelModel().loads(template_path).finish()


def init_session_state() -> None:
    if "version" not in st.session_state:
        st.session_state.version = list(VERSIONS.keys())[0]
    if "manager" not in st.session_state:
        st.session_state.manager = None
    if "active_sheet" not in st.session_state:
        st.session_state.active_sheet = None
    if "recalc_needed" not in st.session_state:
        st.session_state.recalc_needed = False


def load_workbook_for_version(version_label: str) -> WorkbookManager:
    path = get_template_path(version_label)
    manager = WorkbookManager(path)
    model = load_formula_model(str(path), version_label)
    manager.attach_formula_model(model)
    with st.spinner("Calculating workbook formulas..."):
        manager.recalculate()
    return manager


def on_version_change() -> None:
    version = st.session_state.version_selector
    st.session_state.version = version
    st.session_state.manager = load_workbook_for_version(version)
    sheets = ordered_sheets(st.session_state.manager.list_sheets())
    st.session_state.active_sheet = sheets[0]
    st.session_state.recalc_needed = False


def main() -> None:
    st.set_page_config(
        page_title="Subjective Spreadsheet Tool",
        page_icon="📊",
        layout="wide",
        initial_sidebar_state="expanded",
    )

    init_session_state()

    st.markdown(
        """
        <style>
          .block-container { padding-top: 1rem; max-width: 100%; }
          div[data-testid="stSidebar"] { background-color: #f7f7f7; }
          .version-banner {
            background: linear-gradient(90deg, #217346 0%, #185c37 100%);
            color: white;
            padding: 12px 16px;
            border-radius: 8px;
            margin-bottom: 12px;
          }
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.title("Subjective Spreadsheet Tool")
    st.caption("Python Streamlit version of the subjective driveability test spreadsheets")

    version_labels = list(VERSIONS.keys())
    current_index = version_labels.index(st.session_state.version)

    st.selectbox(
        "Select spreadsheet version",
        version_labels,
        index=current_index,
        key="version_selector",
        help="Choose Base (General), BEV, or CVT — same templates as the original Excel files.",
        on_change=on_version_change,
    )

    if st.session_state.manager is None:
        st.session_state.manager = load_workbook_for_version(st.session_state.version)

    manager: WorkbookManager = st.session_state.manager
    sheets = ordered_sheets(manager.list_sheets())

    if st.session_state.active_sheet not in sheets:
        st.session_state.active_sheet = sheets[0]

    st.markdown(
        f"<div class='version-banner'><strong>Active version:</strong> {st.session_state.version}"
        f" &nbsp;|&nbsp; <strong>File:</strong> {manager.book_name}</div>",
        unsafe_allow_html=True,
    )

    with st.sidebar:
        st.header("Sheet Navigation")
        st.markdown("Worksheets match the Excel tab bar order.")

        for group_name, group_sheets in [
            ("Report", ["Test Summary -->", "DRB - Summary", "DRB - Color Chart", "Optional Test - Summary", "Engine - Summary"]),
            ("DRB Tests", ["DRB Testing -->", "Driveaway-sweeps", "Driveaway ESS", "Decel Cstdowns", "Decel OPD", "USS-Manual (Opt.)", "RTITO_AD", "TI_CstSpd ", "TO_CstSpd", "RRL (with HS)", "Garage Shifts", "Kickdowns"]),
            ("Engine Tests", ["ENG Testing -->", "TITO(Opt)", "CstSpd", "Acceleration(Opt)", "Stationary Test", "Engine - Color Chart"]),
            ("Reference", ["Pedal Geo", "Test Details -->", "PV Max", "Acronyms"]),
        ]:
            visible = [s for s in group_sheets if s in sheets]
            if not visible:
                continue
            st.subheader(group_name)
            for sheet_name in visible:
                if st.button(sheet_name, key=f"nav_{sheet_name}", use_container_width=True):
                    st.session_state.active_sheet = sheet_name

        st.divider()
        st.subheader("Recalculate")
        if st.button("Recalculate all sheets", type="primary", use_container_width=True):
            with st.spinner("Recalculating formulas..."):
                manager.recalculate()
            st.session_state.recalc_needed = False
            st.success("Recalculation complete.")

        export_name = manager.book_name.replace(".xlsm", "_export.xlsm")
        st.download_button(
            label="Download Excel (.xlsm)",
            data=manager.export_bytes(),
            file_name=export_name,
            mime="application/vnd.ms-excel.sheet.macroEnabled.12",
            use_container_width=True,
        )

        st.divider()
        st.subheader("Acronym Reference")
        st.markdown(
            "Use shorthand from the **Acronyms** sheet: `nd`, `b`, `c`, `st`, `j`, `sh`, etc. "
            "Prefix with `!` for bad events (e.g. `!b`)."
        )

    active_sheet = st.session_state.active_sheet
    st.subheader(active_sheet)

    bounds = manager.sheet_bounds(active_sheet)
    if bounds:
        min_r, max_r, min_c, max_c = bounds
        st.caption(f"Used range: {min_r}:{max_r} rows, columns {min_c}:{max_c}")
    else:
        st.info("This worksheet is empty in the template — same as the original Excel file.")

    with st.expander("Edit cell values (manual entry sheets)", expanded=True):
        st.markdown(
            "Enter values exactly as you would in Excel. Formula cells are calculated automatically "
            "and shown in the preview below."
        )

        col1, col2, col3 = st.columns([1, 2, 1])
        with col1:
            cell_coord = st.text_input("Cell", value="A1", key=f"cell_coord_{active_sheet}")
        with col2:
            current = manager.get_display_value(active_sheet, cell_coord.upper())
            new_value = st.text_input(
                "Value",
                value=format_display_value(current),
                key=f"cell_value_{active_sheet}",
            )
        with col3:
            apply = st.button("Apply", type="primary", use_container_width=True)

        if apply:
            coord = cell_coord.upper().strip()
            if not manager.is_editable(active_sheet, coord):
                st.error(f"{coord} is a formula cell and cannot be edited directly.")
            else:
                manager.set_cell_value(active_sheet, coord, new_value)
                with st.spinner("Recalculating..."):
                    manager.recalculate()
                st.success(f"Updated {active_sheet}!{coord}")
                st.rerun()

        # Dropdown for PV_Max selector cells
        pv_options = manager.get_named_range_values("PV_Max")
        if pv_options and active_sheet in ("Driveaway-sweeps", "Driveaway ESS", "Decel Cstdowns", "USS-Manual (Opt.)"):
            st.markdown("**Transmission / PV Max selector (cell A1)**")
            current_a1 = format_display_value(manager.get_display_value(active_sheet, "A1"))
            selected_pv = st.selectbox(
                "PV Max variant",
                pv_options,
                index=pv_options.index(current_a1) if current_a1 in pv_options else 0,
                key=f"pvmax_{active_sheet}",
            )
            if st.button("Apply A1 selection", key=f"apply_pv_{active_sheet}"):
                manager.set_cell_value(active_sheet, "A1", selected_pv)
                with st.spinner("Recalculating..."):
                    manager.recalculate()
                st.rerun()

        editable = manager.editable_cells_for_sheet(active_sheet)
        st.caption(f"{len(editable)} editable cells on this sheet")

    st.markdown(render_sheet_css(), unsafe_allow_html=True)
    st.markdown(render_sheet_html(manager, active_sheet), unsafe_allow_html=True)

    with st.expander("Quick edit — batch cell entry"):
        entry_cols = st.multiselect(
            "Columns to edit",
            options=list(range(1, 27)),
            default=[4, 5, 6, 7, 8],
            format_func=lambda c: f"Column {c}",
            key=f"cols_{active_sheet}",
        )
        if bounds and entry_cols:
            min_r, max_r, min_c, max_c = bounds
            from openpyxl.utils import get_column_letter

            with st.form(key=f"batch_form_{active_sheet}"):
                edits: list[tuple[str, Any]] = []
                for r in range(min_r, min(min_r + 35, max_r)):
                    row_coords = []
                    for c in range(min_c, max_c + 1):
                        if c not in entry_cols:
                            continue
                        coord = f"{get_column_letter(c)}{r}"
                        if manager.is_editable(active_sheet, coord):
                            row_coords.append(coord)
                    if not row_coords:
                        continue
                    cols = st.columns(len(row_coords))
                    for i, coord in enumerate(row_coords):
                        with cols[i]:
                            val = format_display_value(manager.get_display_value(active_sheet, coord))
                            edits.append((coord, st.text_input(coord, value=val)))
                submitted = st.form_submit_button("Apply batch edits & recalculate")
                if submitted:
                    for coord, new_val in edits:
                        old_val = format_display_value(manager.get_display_value(active_sheet, coord))
                        if new_val != old_val:
                            manager.set_cell_value(active_sheet, coord, new_val)
                    with st.spinner("Recalculating..."):
                        manager.recalculate()
                    st.rerun()


if __name__ == "__main__":
    main()
