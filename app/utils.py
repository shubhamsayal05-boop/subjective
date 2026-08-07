import re
from typing import Any

from openpyxl.utils import get_column_letter


def color_to_css(color: Any) -> str | None:
    if color is None:
        return None
    rgb = getattr(color, "rgb", None) or getattr(color, "value", None)
    if not rgb or not isinstance(rgb, str):
        return None
    if rgb in ("00000000", "0"):
        return None
    if len(rgb) == 8:
        rgb = rgb[2:]
    if len(rgb) != 6:
        return None
    return f"#{rgb}"


def extract_range_value(raw: Any) -> Any:
    if raw is None:
        return None
    if hasattr(raw, "value"):
        return extract_range_value(raw.value)
    if isinstance(raw, list):
        if not raw:
            return None
        return extract_range_value(raw[0])
    if isinstance(raw, tuple):
        if not raw:
            return None
        return extract_range_value(raw[0])
    if isinstance(raw, (int, float, str, bool)):
        return raw
    return str(raw)


def format_display_value(value: Any) -> str:
    value = extract_range_value(value)
    if value is None:
        return ""
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return f"{value:.4g}"
    return str(value)


def is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def sheet_bounds(ws) -> tuple[int, int, int, int] | None:
    min_r, max_r, min_c, max_c = None, 0, None, 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                r, c = cell.row, cell.column
                min_r = min(min_r or r, r)
                max_r = max(max_r, r)
                min_c = min(min_c or c, c)
                max_c = max(max_c, c)
    if min_r is None:
        return None
    return min_r, max_r, min_c, max_c


def parse_dispatch_key(key: str) -> tuple[str, str] | None:
    match = re.match(r"'\[[^\]]+\]([^']+)'!(.+)$", key)
    if not match:
        return None
    sheet, coord = match.group(1), match.group(2)
    if ":" in coord:
        coord = coord.split(":")[0]
    return sheet, coord


def build_dispatch_key(book_name: str, sheet: str, coord: str) -> str:
    return f"'[{book_name}]{sheet}'!{coord}"


def column_width_px(ws, col_idx: int) -> int:
    letter = get_column_letter(col_idx)
    width = ws.column_dimensions[letter].width
    if width is None:
        width = 8.43
    return max(int(width * 7), 24)


def row_height_px(ws, row_idx: int) -> int:
    height = ws.row_dimensions[row_idx].height
    if height is None:
        height = 15
    return max(int(height * 1.33), 16)
