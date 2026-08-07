import copy
import io
import re
from pathlib import Path
from typing import Any

import formulas
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.utils import (
    build_dispatch_key,
    extract_range_value,
    format_display_value,
    is_formula,
    parse_dispatch_key,
    sheet_bounds,
)


class WorkbookManager:
    def __init__(self, template_path: Path):
        self.template_path = Path(template_path)
        self.book_name = self.template_path.name
        self.wb = load_workbook(self.template_path, keep_vba=True)
        self.formula_model: formulas.ExcelModel | None = None
        self.solution: dict[Any, Any] = {}
        self.user_values: dict[tuple[str, str], Any] = {}
        self._formula_cells: set[tuple[str, str]] = set()
        self._editable_cells: dict[str, list[tuple[str, str]]] = {}
        self._data_nodes: set[str] = set()
        self._scan_cells()

    def _scan_cells(self) -> None:
        for ws in self.wb.worksheets:
            editable: list[tuple[str, str]] = []
            for row in ws.iter_rows():
                for cell in row:
                    coord = cell.coordinate
                    key = (ws.title, coord)
                    if is_formula(cell.value):
                        self._formula_cells.add(key)
                    else:
                        editable.append(key)
            self._editable_cells[ws.title] = editable

    def attach_formula_model(self, model: formulas.ExcelModel) -> None:
        self.formula_model = model
        self._data_nodes = {str(k) for k in model.dsp.data_nodes}

    def get_raw_cell_value(self, sheet: str, coord: str) -> Any:
        key = (sheet, coord)
        if key in self.user_values:
            return self.user_values[key]
        return self.wb[sheet][coord].value

    def get_display_value(self, sheet: str, coord: str) -> Any:
        key = (sheet, coord)
        if key in self._formula_cells or is_formula(self.wb[sheet][coord].value):
            computed = self._solution_lookup(sheet, coord)
            if computed is not None:
                return computed
        if key in self.user_values:
            return self.user_values[key]
        val = self.wb[sheet][coord].value
        if is_formula(val):
            return self._solution_lookup(sheet, coord) or ""
        return val

    def _solution_lookup(self, sheet: str, coord: str) -> Any:
        book_upper = self.book_name.upper()
        sheet_upper = sheet.upper()
        for sol_key, raw in self.solution.items():
            key_str = str(sol_key)
            if book_upper not in key_str.upper():
                continue
            if sheet_upper not in key_str.upper():
                continue
            if key_str.endswith(f"!{coord}") or f"!{coord}:" in key_str:
                return extract_range_value(raw)
            parsed = parse_dispatch_key(key_str)
            if parsed and parsed[0].upper() == sheet_upper and parsed[1] == coord:
                return extract_range_value(raw)
        return None

    def set_cell_value(self, sheet: str, coord: str, value: Any) -> None:
        key = (sheet, coord)
        if key in self._formula_cells:
            return
        if value == "" or value is None:
            self.user_values.pop(key, None)
            self.wb[sheet][coord].value = None
        else:
            original = self.wb[sheet][coord].value
            if isinstance(original, (int, float)) and not isinstance(value, str):
                try:
                    value = float(value) if "." in str(value) else int(value)
                except (ValueError, TypeError):
                    pass
            self.user_values[key] = value
            self.wb[sheet][coord].value = value

    def is_editable(self, sheet: str, coord: str) -> bool:
        return (sheet, coord) not in self._formula_cells

    def is_formula_cell(self, sheet: str, coord: str) -> bool:
        return (sheet, coord) in self._formula_cells

    def editable_cells_for_sheet(self, sheet: str) -> list[tuple[str, str]]:
        return self._editable_cells.get(sheet, [])

    def _build_dispatch_inputs(self) -> dict[str, list[list[Any]]]:
        inputs: dict[str, list[list[Any]]] = {}
        for ws in self.wb.worksheets:
            sheet = ws.title
            for row in ws.iter_rows():
                for cell in row:
                    coord = cell.coordinate
                    key = (sheet, coord)
                    if key in self._formula_cells:
                        continue
                    value = self.user_values.get(key, cell.value)
                    if value is None:
                        continue
                    dispatch_key = build_dispatch_key(self.book_name, sheet, coord)
                    if dispatch_key in self._data_nodes:
                        inputs[dispatch_key] = [[value]]
                    range_key = self._find_range_node(sheet, coord)
                    if range_key and range_key not in inputs:
                        inputs[range_key] = [[value]]
        return inputs

    def _find_range_node(self, sheet: str, coord: str) -> str | None:
        sheet_upper = sheet.upper()
        for node in self._data_nodes:
            parsed = parse_dispatch_key(node)
            if not parsed:
                continue
            if parsed[0].upper() != sheet_upper:
                continue
            if ":" in node:
                range_part = node.split("!")[1].replace("'", "")
                start, end = range_part.split(":")
                if self._coord_in_range(coord, start, end):
                    return node
        return None

    def _coord_in_range(self, coord: str, start: str, end: str) -> bool:
        start_col, start_row = self._split_coord(start)
        end_col, end_row = self._split_coord(end)
        col, row = self._split_coord(coord)
        return start_col <= col <= end_col and start_row <= row <= end_row

    def _split_coord(self, coord: str) -> tuple[int, int]:
        match = re.match(r"([A-Z]+)(\d+)", coord)
        if not match:
            return 0, 0
        col_str, row = match.group(1), int(match.group(2))
        col = 0
        for ch in col_str:
            col = col * 26 + (ord(ch) - ord("A") + 1)
        return col, row

    def recalculate(self) -> None:
        if self.formula_model is None:
            raise RuntimeError("Formula model not loaded")
        inputs = self._build_dispatch_inputs()
        self.solution = self.formula_model.dsp.dispatch(inputs)

    def get_named_range_values(self, name: str) -> list[str]:
        dn = self.wb.defined_names.get(name)
        if dn is None:
            return []
        destinations = list(dn.destinations)
        if not destinations:
            return []
        sheet, ref = destinations[0]
        ws = self.wb[sheet]
        values: list[str] = []
        for row in ws[ref]:
            for cell in row:
                val = self.get_display_value(sheet, cell.coordinate)
                if val is not None and str(val).strip():
                    values.append(format_display_value(val))
        return values

    def export_bytes(self) -> bytes:
        buffer = io.BytesIO()
        export_wb = load_workbook(self.template_path, keep_vba=True)
        for (sheet, coord), value in self.user_values.items():
            cell = export_wb[sheet][coord]
            if not is_formula(cell.value):
                cell.value = value
        export_wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def clone(self) -> "WorkbookManager":
        new = WorkbookManager(self.template_path)
        new.user_values = copy.deepcopy(self.user_values)
        for (sheet, coord), value in new.user_values.items():
            new.wb[sheet][coord].value = value
        new.formula_model = self.formula_model
        new._data_nodes = self._data_nodes
        new.solution = copy.deepcopy(self.solution)
        return new

    def sheet_bounds(self, sheet: str) -> tuple[int, int, int, int] | None:
        return sheet_bounds(self.wb[sheet])

    def list_sheets(self) -> list[str]:
        return self.wb.sheetnames
