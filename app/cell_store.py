import json
import re
from pathlib import Path
from typing import Any

from app.paths import bundled_dir
from app.utils import format_display_value


def load_bundle(key: str) -> dict:
    path = bundled_dir() / f"{key}.json"
    with open(path, encoding="utf-8") as f:
        return json.load(f)


class CellStore:
    def __init__(self, bundle: dict):
        self.bundle = bundle
        self.book_name = bundle["book_name"]
        self.named_ranges = bundle.get("named_ranges", {})
        self.values: dict[tuple[str, str], Any] = {}
        self.formulas: dict[tuple[str, str], str] = {}
        self.user_edits: dict[tuple[str, str], Any] = {}
        self._load_cells()

    def _load_cells(self) -> None:
        for sheet_name, sheet_data in self.bundle["sheets"].items():
            for coord, cell in sheet_data.get("cells", {}).items():
                key = (sheet_name, coord)
                if "formula" in cell:
                    self.formulas[key] = cell["formula"]
                    cached = cell.get("cached", "")
                    self.values[key] = cached
                elif "value" in cell:
                    self.values[key] = cell["value"]

    def list_sheets(self) -> list[str]:
        return list(self.bundle["sheets"].keys())

    def sheet_meta(self, sheet: str) -> dict:
        return self.bundle["sheets"].get(sheet, {})

    def get(self, sheet: str, coord: str) -> Any:
        key = (sheet, coord)
        if key in self.user_edits:
            return self.user_edits[key]
        return self.values.get(key, "")

    def set_user_value(self, sheet: str, coord: str, value: Any) -> None:
        key = (sheet, coord)
        if key in self.formulas:
            return
        if value == "" or value is None:
            self.user_edits.pop(key, None)
            if key in self.values:
                self.values[key] = ""
        else:
            self.user_edits[key] = value
            self.values[key] = value

    def set_computed(self, sheet: str, coord: str, value: Any) -> None:
        self.values[(sheet, coord)] = value

    def is_editable(self, sheet: str, coord: str) -> bool:
        return (sheet, coord) not in self.formulas

    def is_formula(self, sheet: str, coord: str) -> bool:
        return (sheet, coord) in self.formulas

    def display_value(self, sheet: str, coord: str) -> str:
        return format_display_value(self.get(sheet, coord))

    def editable_count(self, sheet: str) -> int:
        sheet_data = self.bundle["sheets"].get(sheet, {})
        count = 0
        for coord, cell in sheet_data.get("cells", {}).items():
            if "formula" not in cell:
                count += 1
        return count

    def get_named_range_values(self, name: str) -> list[str]:
        ref = self.named_ranges.get(name)
        if not ref:
            return []
        sheet, range_ref = ref.split("!", 1)
        values = []
        for coord in self._expand_range(sheet, range_ref):
            val = self.display_value(sheet, coord)
            if val:
                values.append(val)
        return values

    def range_values(self, sheet: str, start: str, end: str) -> list[Any]:
        coords = self._expand_range(sheet, f"{start}:{end}")
        return [self.get(sheet, c) for c in coords]

    def range_values_2d(self, sheet: str, start: str, end: str) -> list[list[Any]]:
        start = re.sub(r"\$", "", start)
        end = re.sub(r"\$", "", end)
        sc, sr = _split_coord(start)
        ec, er = _split_coord(end)
        rows = []
        for r in range(sr, er + 1):
            row_vals = []
            for c in range(sc, ec + 1):
                row_vals.append(self.get(sheet, _coord_str(c, r)))
            rows.append(row_vals)
        return rows

    def _expand_range(self, sheet: str, range_ref: str) -> list[str]:
        if ":" in range_ref:
            start, end = range_ref.split(":")
            return _coords_in_range(start, end)
        return [re.sub(r"\$", "", range_ref)]

    def export_xlsx_bytes(self) -> bytes:
        import io

        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        for sheet_name in self.list_sheets():
            meta = self.sheet_meta(sheet_name)
            ws = wb.create_sheet(sheet_name)
            bounds = meta.get("bounds")
            if not bounds:
                continue
            min_r, max_r, min_c, max_c = bounds
            from openpyxl.utils import get_column_letter

            for r in range(min_r, max_r + 1):
                for c in range(min_c, max_c + 1):
                    coord = f"{get_column_letter(c)}{r}"
                    val = self.get(sheet_name, coord)
                    if val != "" and val is not None:
                        ws[coord].value = _to_export_value(val)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.getvalue()


def _to_export_value(val: Any) -> Any:
    if isinstance(val, (int, float)):
        return val
    text = str(val)
    try:
        if "." in text:
            return float(text)
        return int(text)
    except ValueError:
        return text


def _coords_in_range(start: str, end: str) -> list[str]:
    start = re.sub(r"\$", "", start)
    end = re.sub(r"\$", "", end)
    sc, sr = _split_coord(start)
    ec, er = _split_coord(end)
    coords = []
    for r in range(sr, er + 1):
        for c in range(sc, ec + 1):
            coords.append(_coord_str(c, r))
    return coords


def _split_coord(coord: str) -> tuple[int, int]:
    m = re.match(r"([A-Z]+)(\d+)", coord)
    if not m:
        return 1, 1
    col = 0
    for ch in m.group(1):
        col = col * 26 + (ord(ch) - ord("A") + 1)
    return col, int(m.group(2))


def _coord_str(col: int, row: int) -> str:
    s = ""
    while col > 0:
        col, rem = divmod(col - 1, 26)
        s = chr(65 + rem) + s
    return f"{s}{row}"
